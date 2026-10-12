#!/usr/bin/env python3
"""Inspect the official WHO 2007 BMI-for-age expanded tables (structure probe)."""
import openpyxl, sys

for f in ["bmi-boys-z.xlsx", "bmi-girls-z.xlsx", "bmi-boys-perc.xlsx", "bmi-girls-perc.xlsx"]:
    wb = openpyxl.load_workbook(f"/tmp/{f}", data_only=True)
    print(f"\n=== {f} ===")
    for ws in wb.worksheets:
        print(f"  sheet '{ws.title}': {ws.max_row} rows x {ws.max_column} cols")
        # print first 4 rows, first 12 cols
        for r in range(1, min(5, ws.max_row + 1)):
            vals = []
            for c in range(1, min(13, ws.max_column + 1)):
                v = ws.cell(row=r, column=c).value
                vals.append(str(v)[:12] if v is not None else "·")
            print(f"    r{r}: {vals}")
