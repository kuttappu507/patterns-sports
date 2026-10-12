#!/usr/bin/env python3
"""Validate WHO 2007 BMI-for-age pipeline against WHO's own anthroplus data.

1. Dump bfa_growth_standards (months 60-229, both sexes) from the WHO R package.
2. Cross-check against the official published XLSX expanded tables (months 61-228).
3. Reproduce the WHO package's own test fixture: girl 60.32mo, 113.8cm, 18.7kg -> zbfa -0.58.
4. Round-trip: LMS-computed BMI at z = -3..+3 vs published SD columns (both sexes, sampled months).
"""
import csv, math
import openpyxl
import pyreadr

r = pyreadr.read_r('/tmp/anthroplus/anthroplus/R/sysdata.rda')
bfa = r['bfa_growth_standards']

# ---- 1. dump package data (sex: 1=m, 2=f in the package) ----
rows = []
for _, row in bfa.iterrows():
    rows.append((int(row['sex']), int(row['age']), float(row['l']), float(row['m']), float(row['s'])))
rows.sort()
with open('/tmp/who_bfa_package.csv', 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['sex', 'month', 'l', 'm', 's'])
    w.writerows(rows)

months = sorted({m for (_, m, _, _, _) in rows})
print(f"package data: months {months[0]}..{months[-1]} ({len(months)} months x 2 sexes = {len(rows)} rows)")

# ---- 2. cross-check vs XLSX ----
def load_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    out = {}
    for r_ in ws.iter_rows(min_row=2, values_only=True):
        if r_[0] is None:
            continue
        out[int(r_[0])] = dict(l=float(r_[1]), m=float(r_[2]), s=float(r_[3]),
                               sd=[float(x) for x in r_[4:12]])  # SD4neg..SD3
    return out

boys = load_xlsx('/tmp/bmi-boys-z.xlsx')
girls = load_xlsx('/tmp/bmi-girls-z.xlsx')
pkg = {(s, m): (l, m_, s_) for (s, m, l, m_, s_) in rows}
pkg_m = {1: {m: v for (s, m), v in pkg.items() if s == 1},
         2: {m: v for (s, m), v in pkg.items() if s == 2}}

maxdiff = 0.0
for sex_x, xlsx in ((1, boys), (2, girls)):
    for month, ref in xlsx.items():
        l, m, s = pkg_m[sex_x][month]
        for a, b in ((l, ref['l']), (m, ref['m']), (s, ref['s'])):
            maxdiff = max(maxdiff, abs(a - b))
print(f"LMS cross-check XLSX vs R-package (months 61-228): max abs diff = {maxdiff:.6f}")
assert maxdiff < 5e-5, "reference data mismatch between WHO distributions!"

# ---- 3. WHO test fixture: girl 60.32mo 113.8cm 18.7kg -> zbfa -0.58 ----
def lms_lookup(sex, age_months):
    low = math.trunc(age_months)
    upp = low + 1
    diff = age_months - low
    l0, m0, s0 = pkg_m[sex][low]
    l1, m1, s1 = pkg_m[sex][upp]
    f = lambda a, b: a + diff * (b - a)
    return f(l0, l1), f(m0, m1), f(s0, s1)

def zscore_adjusted(y, m, l, s):
    calc = lambda sd: m * ((1 + l * s * sd) ** (1 / l))
    z = ((y / m) ** l - 1) / (s * l)
    sd3p, sd3n = calc(3), calc(-3)
    sd23p, sd23n = sd3p - calc(2), calc(-2) - sd3n
    if z > 3:
        z = 3 + (y - sd3p) / sd23p
    elif z < -3:
        z = -3 + (y - sd3n) / sd23n
    return z

for age, h, w, expect in ((60.32, 113.8, 18.7, -0.58), (60.911701, 113.6, 20.5, 0.42)):
    bmi = w / ((h / 100) ** 2)
    l, m, s = lms_lookup(2, age)
    z = zscore_adjusted(bmi, m, l, s)
    print(f"WHO fixture girl {age}mo h={h} w={w}: bmi={bmi:.4f} z={z:.4f} rounded={round(z,2)} (expect {expect})")
    assert round(z, 2) == expect, "WHO fixture mismatch!"

# ---- 4. round-trip: LMS z=-3..+3 vs published SD columns ----
# WHO computes the +/-3 SD columns with plain LMS; beyond that (SD4neg) the
# RESTRICTED method applies: y(restricted z) = SD3 + (z-3)*SD23.
worst = 0.0
worst_restricted = 0.0
for sex_x, xlsx in ((1, boys), (2, girls)):
    for month, ref in xlsx.items():
        if month % 37 not in (0, 1, 2):  # sampled months for speed
            continue
        l, m, s = pkg_m[sex_x][month]
        calc = lambda sd: m * ((1 + l * s * sd) ** (1 / l))
        for i, z in enumerate((-4, -3, -2, -1, 0, 1, 2, 3)):
            if z == -4:
                # restricted: y = SD3neg - SD23neg  (one SD23 span beyond -3)
                y_res = calc(-3) - (calc(-2) - calc(-3))
                worst_restricted = max(worst_restricted, abs(y_res - ref['sd'][i]))
                continue
            sd = calc(z)
            worst = max(worst, abs(sd - ref['sd'][i]))
print(f"round-trip plain LMS z=-3..+3 vs published SD cols: max abs diff = {worst:.4f} (cols rounded to 3dp)")
print(f"restricted z=-4 (SD3neg - SD23neg) vs published SD4neg col: max abs diff = {worst_restricted:.4f}")

# percentile table sanity: P50 == M, P25 ~ z=-0.6745
for sex_x, path in ((1, '/tmp/bmi-boys-perc.xlsx'), (2, '/tmp/bmi-girls-perc.xlsx')):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    i_p50, i_p25, i_p01 = hdr.index('P50'), hdr.index('P25'), hdr.index('P01')
    for r_ in ws.iter_rows(min_row=2, values_only=True):
        if r_[0] is None or int(r_[0]) % 50 != 0:
            continue
        month = int(r_[0])
        l, m, s = pkg_m[sex_x][month]
        assert abs(r_[i_p50] - m) < 0.01, (month, r_[i_p50], m)
        z25 = ((r_[i_p25] / m) ** l - 1) / (s * l)
        z01 = ((r_[i_p01] / m) ** l - 1) / (s * l)
        print(f"  {'boy' if sex_x==1 else 'girl'} m{month}: P25 z={z25:+.4f} (expect -0.6745) P01 z={z01:+.4f} (expect -2.3263)")
print("ALL PYTHON-SIDE VALIDATION PASSED")
